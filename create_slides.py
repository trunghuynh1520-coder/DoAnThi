from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

OUT = 'Thuyet_trinh_Tuan_06_Sprint_2.pptx'
prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)

NAVY = RGBColor(22, 38, 52)
INK = RGBColor(31, 43, 52)
MUTED = RGBColor(100, 116, 124)
CREAM = RGBColor(248, 245, 237)
PAPER = RGBColor(255, 253, 248)
TEAL = RGBColor(30, 126, 119)
CORAL = RGBColor(212, 105, 78)
GOLD = RGBColor(221, 174, 85)
PALE = RGBColor(231, 237, 229)
WHITE = RGBColor(255, 255, 255)

def textbox(slide, x, y, w, h, text, size=16, color=INK, bold=False, font='Aptos', align=None):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.clear(); tf.word_wrap = True
    tf.margin_left = Inches(.02); tf.margin_right = Inches(.02)
    tf.margin_top = Inches(.02); tf.margin_bottom = Inches(.02)
    p = tf.paragraphs[0]
    if align is not None: p.alignment = align
    r = p.add_run(); r.text = text
    r.font.name = font; r.font.size = Pt(size); r.font.bold = bold; r.font.color.rgb = color
    return box

def rect(slide, x, y, w, h, fill, radius=False, line=None):
    shp = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE,
                                 Inches(x), Inches(y), Inches(w), Inches(h))
    shp.fill.solid(); shp.fill.fore_color.rgb = fill
    shp.line.fill.background() if line is None else None
    if line is not None: shp.line.color.rgb = line
    return shp

def base(kicker, title, number):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    s.background.fill.solid(); s.background.fill.fore_color.rgb = CREAM
    rect(s, 0, 0, .13, 7.5, TEAL)
    textbox(s, .55, .32, 8, .3, kicker.upper(), 10, TEAL, True)
    textbox(s, .55, .68, 11.8, .62, title, 27, NAVY, True)
    rect(s, .55, 1.42, 12.2, .018, RGBColor(215, 219, 211))
    textbox(s, .6, 7.13, 9, .2, 'ĐỀ TÀI 11  •  NHÀ HÀNG MÓN HUẾ', 9, MUTED, True)
    textbox(s, 12.1, 7.1, .55, .24, f'{number:02d} / 05', 9, MUTED, True, align=PP_ALIGN.RIGHT)
    return s

def card(s, x, y, w, h, title, body, accent=TEAL, body_size=15):
    rect(s, x, y, w, h, PAPER, True)
    rect(s, x, y, .07, h, accent)
    textbox(s, x+.22, y+.19, w-.42, .35, title, 16, NAVY, True)
    textbox(s, x+.22, y+.66, w-.42, h-.78, body, body_size, INK)

# 1 — Cover and domain mapping
s = prs.slides.add_slide(prs.slide_layouts[6])
s.background.fill.solid(); s.background.fill.fore_color.rgb = NAVY
rect(s, 0, 0, .16, 7.5, TEAL)
textbox(s, .8, .65, 5, .35, 'TUẦN 06  /  SPRINT 2', 12, GOLD, True)
textbox(s, .8, 1.25, 7.1, 1.55, 'NHÀ HÀNG MÓN HUẾ\nĐẶT BÀN & GỌI MÓN', 31, WHITE, True)
textbox(s, .8, 3.15, 6.1, .85, 'Từ mô hình eShop đến luồng phục vụ tại bàn', 19, RGBColor(210, 221, 218))
rect(s, 8.0, .75, 4.55, 5.85, RGBColor(31, 57, 68), True)
textbox(s, 8.35, 1.05, 3.8, .35, 'ÁNH XẠ MIỀN NGHIỆP VỤ', 11, GOLD, True)
mapping = [('Product', 'MenuItem'), ('Category', 'MenuCategory'), ('ShoppingCart', 'ScopedOrderState'), ('Order / OrderLine', 'TableOrder / TableOrderLine')]
for i,(a,b) in enumerate(mapping):
    yy=1.72+i*1.02
    textbox(s, 8.35, yy, 1.58, .33, a, 13, RGBColor(200,215,212), True)
    textbox(s, 10.05, yy, .3, .33, '→', 16, GOLD, True, align=PP_ALIGN.CENTER)
    textbox(s, 10.42, yy, 1.75, .58, b, 13, WHITE, True)
textbox(s, .8, 5.7, 6.7, .35, 'Thành viên: ______________________________', 12, RGBColor(179, 197, 193))
textbox(s, .8, 7.08, 8, .2, 'ĐỀ TÀI 11  •  ĐỘ KHÓ ★★', 9, RGBColor(179,197,193), True)
textbox(s, 12.1, 7.08, .55, .2, '01 / 05', 9, RGBColor(179,197,193), True, align=PP_ALIGN.RIGHT)

# 2 — Architecture / database
s=base('01  /  THIẾT KẾ HỆ THỐNG','Kiến trúc 4 tầng, dữ liệu tập trung SQL Server',2)
layers=[('WebApp','Blazor InteractiveServer\nUI • DI Composition Root',TEAL),
        ('UseCases','Luồng ứng dụng\nRepository interfaces',RGBColor(64,145,133)),
        ('CoreBusiness','Entities & luật nghiệp vụ\nđộc lập giao diện',CORAL),
        ('SQL Plugin','Dapper repositories\nSQL Server',GOLD)]
for i,(name,desc,c) in enumerate(layers):
    x=.72+i*3.12
    rect(s,x,1.9,2.75,1.68,PAPER,True)
    rect(s,x,1.9,2.75,.12,c)
    textbox(s,x+.18,2.18,2.4,.36,name,17,NAVY,True,align=PP_ALIGN.CENTER)
    textbox(s,x+.18,2.67,2.4,.67,desc,13,MUTED,False,align=PP_ALIGN.CENTER)
    if i<3: textbox(s,x+2.78,2.48,.34,.35,'→',17,TEAL,True,align=PP_ALIGN.CENTER)
textbox(s,.75,3.92,5.3,.34,'MŨI TÊN PHỤ THUỘC HƯỚNG VÀO USE CASE / DOMAIN',10,TEAL,True)
card(s,.72,4.45,3.82,1.8,'6 bảng dữ liệu','MenuCategory • MenuItem • DiningTable\nReservation • TableOrder • TableOrderLine',TEAL,13)
card(s,4.75,4.45,3.82,1.8,'Ghi phiếu theo transaction','Header và các dòng món được ghi trong\ncùng transaction; có commit / rollback.',CORAL,13)
card(s,8.78,4.45,3.82,1.8,'Truy vấn an toàn','Dapper với tham số SQL (@param);\nkhông ghép dữ liệu đầu vào vào câu lệnh.',GOLD,13)
textbox(s,.78,6.52,11.7,.28,'Công nghệ theo project: .NET 9 • Blazor • Dapper • SQL Server',11,MUTED,True)

# 3 — Rules
s=base('02  /  TRỌNG TÂM VẤN ĐÁP','Hai luật nghiệp vụ được đặt trong CoreBusiness',3)
card(s,.72,1.85,5.9,3.95,'01  ĐẶT BÀN','• Khách không vượt sức chứa của bàn.\n\n• Hai đặt chỗ của cùng bàn cách nhau ít nhất 2 giờ.\n\n• Kiểm tra tại Reservation.IsValidReservation();\n  CreateReservationUseCase gọi trước khi lưu.',TEAL,16)
card(s,6.78,1.85,5.84,3.95,'02  GỌI MÓN & ĐÓNG BILL','• Chỉ thêm món khi bàn đang được sử dụng.\n\n• Tổng hóa đơn = tổng dòng món + phí phục vụ 10%.\n\n• Hóa đơn đã đóng không nhận thêm món.\n\n• Nằm trong TableOrder.AddLineItem() / CloseOrder().',CORAL,16)
rect(s,.72,6.12,11.9,.62,PALE,True)
textbox(s,.98,6.27,11.4,.28,'UI → UseCase → Entity kiểm tra luật → Repository lưu dữ liệu',14,NAVY,True,align=PP_ALIGN.CENTER)

# 4 — end-to-end progress
s=base('03  /  TIẾN ĐỘ SPRINT','Luồng nghiệp vụ và phần quản trị đã có trong mã nguồn',4)
steps=[('01','ĐẶT BÀN','Form nhập thông tin\nkiểm tra dữ liệu'),('02','CHỌN MÓN','Lọc thực đơn\nthêm món vào phiếu'),('03','LƯU PHIẾU','Repository Dapper\ntransaction header + line'),('04','THANH TOÁN','Đóng bill, tính phí\nxuất phiếu chi tiết')]
for i,(n,t,b) in enumerate(steps):
    x=.72+i*3.12
    rect(s,x,1.95,2.7,1.55,PAPER,True)
    textbox(s,x+.18,2.12,.5,.3,n,12,TEAL,True)
    textbox(s,x+.18,2.52,2.36,.3,t,15,NAVY,True)
    textbox(s,x+.18,2.94,2.36,.45,b,12,MUTED)
    if i<3: textbox(s,x+2.73,2.49,.33,.32,'›',22,CORAL,True,align=PP_ALIGN.CENTER)
card(s,.72,4.05,3.82,1.72,'Quản trị bàn','/admin/tables • yêu cầu vai trò Admin\nCookie Authentication',TEAL,13)
card(s,4.75,4.05,3.82,1.72,'Giao diện dùng lại','MenuItemCard • TableStatusBadge\nOrderSummaryPanel',CORAL,13)
card(s,8.78,4.05,3.82,1.72,'Xuất phiếu','ExportBillUseCase đọc dữ liệu\nphiếu và tạo nội dung hóa đơn',GOLD,13)
textbox(s,.8,6.2,11.6,.65,'Khi demo: mở SQL Server và chạy thử luồng đặt bàn → gọi món → đóng bill để xác nhận cấu hình kết nối và dữ liệu thực tế.',12,MUTED)

# 5 — evidence and next steps
s=base('04  /  MINH CHỨNG & DEMO','Hồ sơ bảo vệ và các bước chốt trước buổi trình bày',5)
card(s,.72,1.85,5.8,2.05,'TÀI LIỆU ĐÃ CÓ','GRADING.md • đối chiếu rubric K1–K5\nDocs/ • use case, class, sequence diagrams',TEAL,15)
card(s,6.68,1.85,5.95,2.05,'MÃ NGUỒN','4 project theo Clean Architecture\nSQL schema tại Plugins.DataStore.SQL/schema.sql',CORAL,15)
textbox(s,.82,4.35,4.2,.35,'CHECKLIST DEMO',12,TEAL,True)
check=[('01','Khởi động ứng dụng và xác nhận kết nối SQL Server'),('02','Thử đặt bàn hợp lệ và trường hợp bị trùng / quá sức chứa'),('03','Gọi món, đóng hóa đơn và kiểm tra phí phục vụ 10%'),('04','Đăng nhập Admin, mở trang bàn và xuất phiếu thanh toán')]
for i,(n,t) in enumerate(check):
    yy=4.82+i*.43
    textbox(s,.85,yy,.42,.26,n,10,CORAL,True)
    textbox(s,1.35,yy,10.9,.32,t,13,INK)
rect(s,.72,6.67,11.9,.34,NAVY,True)
textbox(s,.9,6.7,11.5,.22,'GitHub: github.com/trunghuynh1520-coder/DoAnThi',10,WHITE,True,align=PP_ALIGN.CENTER)

prs.save(OUT)
print(OUT)
